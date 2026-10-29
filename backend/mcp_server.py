"""Servidor MCP "retail-bi": solo lectura sobre la tabla `ventas` de Supabase.

Excluye siempre cancelaciones y price<=0, salvo en tasa_cancelacion (que sí
necesita ver las cancelaciones para calcular la tasa). Las herramientas
`generar_informe` y `generar_presentacion` son la única excepción de
escritura: guardan un archivo de salida (informe/presentación) en `outputs/`
con datos ya consultados; nunca modifican la tabla `ventas` ni ningún dato
de origen.
"""
import re
import sys
from datetime import datetime
from pathlib import Path

from mcp.server.mcpserver import MCPServer
from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import consultar

mcp = MCPServer("retail-bi")

RAIZ = Path(__file__).resolve().parent.parent
OUTPUTS_DIR = RAIZ / "outputs"
OUTPUTS_DIR.mkdir(exist_ok=True)


def _nombre_archivo(titulo: str, extension: str) -> Path:
    """Convierte un título en un nombre de archivo seguro dentro de outputs/
    (sin rutas ni caracteres que permitan escapar de esa carpeta)."""
    slug = re.sub(r"[^a-z0-9]+", "-", titulo.lower()).strip("-") or "informe"
    marca_tiempo = datetime.now().strftime("%Y%m%d-%H%M%S")
    return OUTPUTS_DIR / f"{slug}-{marca_tiempo}.{extension}"

FILTRO_VENTAS_VALIDAS = "NOT es_cancelacion AND price > 0"
# El dataset mezcla códigos que no son productos reales con los stock_code de
# artículos: gastos de envío (POST, DOT), comisiones (AMAZONFEE), cargos bancarios
# (BANK CHARGES), descuentos (D), ajustes manuales (M, ADJUST), pruebas (TEST...),
# tarjetas regalo (gift_...), etc. Todos distorsionan el ABC y el ranking de
# cancelación (ej. "Discount" aparecía primero en tasa_cancelacion). Los productos
# reales del catálogo siempre tienen stock_code de 5 cifras al inicio.
FILTRO_PRODUCTO_REAL = "stock_code ~ '^[0-9]{5}'"


@mcp.tool()
def kpi_resumen() -> dict:
    """Devuelve el resumen global de KPI: facturación total, número de productos,
    número de productos clase A y el % de facturación que representa la clase A.
    Útil como primer paso para dar contexto general antes de responder preguntas
    más específicas sobre riesgo o cancelaciones."""
    try:
        filas = consultar(f"""
            WITH facturacion_producto AS (
                SELECT stock_code, SUM(importe) AS facturacion
                FROM ventas
                WHERE {FILTRO_VENTAS_VALIDAS} AND {FILTRO_PRODUCTO_REAL}
                GROUP BY stock_code
            ),
            acumulado AS (
                SELECT
                    stock_code, facturacion,
                    SUM(facturacion) OVER (ORDER BY facturacion DESC) AS facturacion_acumulada,
                    SUM(facturacion) OVER () AS facturacion_total
                FROM facturacion_producto
            )
            SELECT
                COUNT(*) AS n_productos,
                SUM(facturacion) AS facturacion_total,
                COUNT(*) FILTER (WHERE (facturacion_acumulada - facturacion) / facturacion_total < 0.8) AS n_productos_a,
                SUM(facturacion) FILTER (WHERE (facturacion_acumulada - facturacion) / facturacion_total < 0.8) AS facturacion_a
            FROM acumulado;
        """)
        fila = filas[0]
        facturacion_total = float(fila["facturacion_total"])
        facturacion_a = float(fila["facturacion_a"])
        return {
            "facturacion_total": round(facturacion_total, 2),
            "n_productos": fila["n_productos"],
            "n_productos_a": fila["n_productos_a"],
            "pct_facturacion_a": round(facturacion_a / facturacion_total, 4),
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def productos_en_riesgo(umbral_caida: float = 0.30, limite: int = 10) -> list[dict]:
    """Devuelve los productos de clase A (los que generan el 80% de la facturación)
    cuyas ventas de sep-nov 2011 cayeron más de `umbral_caida` (ej. 0.30 = 30%)
    frente al mismo periodo sep-nov 2010. Ordenados de mayor a menor caída.
    Úsala para responder preguntas tipo "¿qué productos importantes están cayendo?".

    Importante sobre `variacion_pct`: un valor de -100% (-1.0) significa que el
    producto tuvo 0 de ventas en sep-nov 2011, lo que puede indicar que dejó de
    venderse/descatalogarse, no solo una caída de demanda. Un valor entre -30% y
    -99% indica una caída parcial, con ventas aún activas pero en descenso. Como
    el resultado se ordena de mayor a menor caída, con `limite` bajo (ej. 10) es
    probable que solo veas productos en -100%, ya que suele haber varias decenas
    empatados en ese valor antes de llegar a las caídas parciales. Al presentar
    estos datos, acláralo brevemente para que no se confunda "sin ventas" con
    "cayendo poco a poco"."""
    try:
        filas = consultar(f"""
            WITH facturacion_producto AS (
                SELECT stock_code, SUM(importe) AS facturacion
                FROM ventas
                WHERE {FILTRO_VENTAS_VALIDAS} AND {FILTRO_PRODUCTO_REAL}
                GROUP BY stock_code
            ),
            acumulado AS (
                SELECT
                    stock_code, facturacion,
                    SUM(facturacion) OVER (ORDER BY facturacion DESC) AS facturacion_acumulada,
                    SUM(facturacion) OVER () AS facturacion_total
                FROM facturacion_producto
            ),
            clase_a AS (
                SELECT stock_code
                FROM acumulado
                WHERE (facturacion_acumulada - facturacion) / facturacion_total < 0.8
            ),
            ventas_2011 AS (
                SELECT stock_code, SUM(importe) AS facturacion_2011
                FROM ventas
                WHERE {FILTRO_VENTAS_VALIDAS} AND {FILTRO_PRODUCTO_REAL}
                  AND invoice_date >= '2011-09-01' AND invoice_date < '2011-12-01'
                GROUP BY stock_code
            ),
            ventas_2010 AS (
                SELECT stock_code, SUM(importe) AS facturacion_2010
                FROM ventas
                WHERE {FILTRO_VENTAS_VALIDAS} AND {FILTRO_PRODUCTO_REAL}
                  AND invoice_date >= '2010-09-01' AND invoice_date < '2010-12-01'
                GROUP BY stock_code
            ),
            -- algunos stock_code tienen más de una descripción en el dataset; nos
            -- quedamos con la más frecuente para no duplicar filas en el resultado
            descripcion AS (
                SELECT stock_code, description FROM (
                    SELECT stock_code, description, COUNT(*) AS n,
                           ROW_NUMBER() OVER (PARTITION BY stock_code ORDER BY COUNT(*) DESC) AS rn
                    FROM ventas
                    WHERE description IS NOT NULL AND description != ''
                    GROUP BY stock_code, description
                ) t WHERE rn = 1
            )
            SELECT
                a.stock_code,
                d.description,
                v10.facturacion_2010 AS ventas_2010,
                COALESCE(v11.facturacion_2011, 0) AS ventas_2011,
                (COALESCE(v11.facturacion_2011, 0) - v10.facturacion_2010) / v10.facturacion_2010 AS variacion_pct
            FROM clase_a a
            JOIN ventas_2010 v10 ON v10.stock_code = a.stock_code
            LEFT JOIN ventas_2011 v11 ON v11.stock_code = a.stock_code
            LEFT JOIN descripcion d ON d.stock_code = a.stock_code
            WHERE v10.facturacion_2010 > 0
              AND (COALESCE(v11.facturacion_2011, 0) - v10.facturacion_2010) / v10.facturacion_2010 < -%s
            ORDER BY variacion_pct ASC
            LIMIT %s;
        """, (umbral_caida, limite))
        return [
            {
                "stock_code": f["stock_code"],
                "description": f["description"],
                "ventas_2010": round(float(f["ventas_2010"]), 2),
                "ventas_2011": round(float(f["ventas_2011"]), 2),
                "variacion_pct": round(float(f["variacion_pct"]), 4),
            }
            for f in filas
        ]
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def tasa_cancelacion(min_unidades: int = 100, limite: int = 10) -> list[dict]:
    """Devuelve los productos con mayor tasa de cancelación (unidades canceladas /
    unidades vendidas), considerando solo productos con al menos `min_unidades`
    unidades vendidas. Útil para preguntas tipo "¿qué productos tienen más
    cancelaciones / problemas de calidad?"."""
    try:
        filas = consultar(f"""
            WITH vendidas AS (
                SELECT stock_code, SUM(quantity) AS unidades_vendidas
                FROM ventas
                WHERE {FILTRO_VENTAS_VALIDAS} AND {FILTRO_PRODUCTO_REAL}
                GROUP BY stock_code
                HAVING SUM(quantity) >= %s
            ),
            canceladas AS (
                SELECT stock_code, SUM(ABS(quantity)) AS unidades_canceladas
                FROM ventas
                WHERE es_cancelacion AND {FILTRO_PRODUCTO_REAL}
                GROUP BY stock_code
            ),
            -- algunos stock_code tienen más de una descripción en el dataset; nos
            -- quedamos con la más frecuente para no duplicar filas en el resultado
            descripcion AS (
                SELECT stock_code, description FROM (
                    SELECT stock_code, description, COUNT(*) AS n,
                           ROW_NUMBER() OVER (PARTITION BY stock_code ORDER BY COUNT(*) DESC) AS rn
                    FROM ventas
                    WHERE description IS NOT NULL AND description != ''
                    GROUP BY stock_code, description
                ) t WHERE rn = 1
            )
            SELECT
                v.stock_code,
                d.description,
                v.unidades_vendidas,
                COALESCE(c.unidades_canceladas, 0) AS unidades_canceladas,
                COALESCE(c.unidades_canceladas, 0)::numeric / v.unidades_vendidas AS tasa
            FROM vendidas v
            LEFT JOIN canceladas c ON c.stock_code = v.stock_code
            LEFT JOIN descripcion d ON d.stock_code = v.stock_code
            ORDER BY tasa DESC
            LIMIT %s;
        """, (min_unidades, limite))
        return [
            {
                "stock_code": f["stock_code"],
                "description": f["description"],
                "unidades_vendidas": f["unidades_vendidas"],
                "unidades_canceladas": f["unidades_canceladas"],
                "tasa": round(float(f["tasa"]), 4),
            }
            for f in filas
        ]
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def ventas_producto(stock_code: str) -> list[dict]:
    """Devuelve la serie mensual de unidades e importe vendidos de un producto
    (identificado por `stock_code`), excluyendo cancelaciones y price<=0. Útil
    para ver la evolución histórica de un producto concreto mes a mes."""
    try:
        filas = consultar(f"""
            SELECT
                to_char(date_trunc('month', invoice_date), 'YYYY-MM') AS mes,
                SUM(quantity) AS unidades,
                SUM(importe) AS importe
            FROM ventas
            WHERE {FILTRO_VENTAS_VALIDAS} AND {FILTRO_PRODUCTO_REAL} AND stock_code = %s
            GROUP BY date_trunc('month', invoice_date)
            ORDER BY date_trunc('month', invoice_date);
        """, (stock_code,))
        return [
            {
                "mes": f["mes"],
                "unidades": f["unidades"],
                "importe": round(float(f["importe"]), 2),
            }
            for f in filas
        ]
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def generar_informe(titulo: str, contenido_markdown: str) -> dict:
    """Guarda un informe en formato Markdown (.md) dentro de la carpeta `outputs/`,
    con el contenido que le pases (ya redactado, con los datos de las otras
    herramientas). Úsala cuando el usuario pida "guárdalo en un archivo",
    "descárgalo" o "dame un informe/documento" con texto, tablas o listas.
    No sirve para PowerPoint: para eso usa `generar_presentacion`.
    No modifica ningún dato de `ventas`, solo crea un archivo nuevo de salida."""
    try:
        ruta = _nombre_archivo(titulo, "md")
        ruta.write_text(f"# {titulo}\n\n{contenido_markdown}\n", encoding="utf-8")
        return {
            "archivo": str(ruta.relative_to(RAIZ)),
            "mensaje": f"Informe guardado en {ruta.relative_to(RAIZ)}",
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def generar_presentacion(titulo: str = "Agente de BI - Productos clave en riesgo") -> dict:
    """Genera una presentación PowerPoint (.pptx) en `outputs/` con los datos
    REALES ya consultados en Supabase: resumen de KPI, productos clase A en
    riesgo y productos con más cancelaciones. Úsala cuando el usuario pida
    "una presentación", "un PowerPoint" o "un .pptx". No necesita argumentos
    de datos: ella misma vuelve a consultar las herramientas internamente,
    así que los números siempre salen de la base de datos, nunca inventados.
    No modifica ningún dato de `ventas`, solo crea un archivo nuevo de salida."""
    try:
        kpi = kpi_resumen()
        riesgo = productos_en_riesgo(0.30, 10)
        cancelacion = tasa_cancelacion(100, 10)

        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        # Diapositiva 1: título
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = titulo
        slide.placeholders[1].text = "Datos reales: Online Retail II (Supabase)"

        # Diapositiva 2: resumen de KPI
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Resumen de KPI"
        cuerpo = slide.placeholders[1].text_frame
        if isinstance(kpi, dict) and "error" not in kpi:
            cuerpo.text = f"Facturación total: {kpi['facturacion_total']:,.2f}"
            cuerpo.add_paragraph().text = f"Productos totales: {kpi['n_productos']}"
            cuerpo.add_paragraph().text = f"Productos clase A: {kpi['n_productos_a']}"
            cuerpo.add_paragraph().text = f"% facturación clase A: {kpi['pct_facturacion_a']:.2%}"
        else:
            cuerpo.text = f"No se pudo calcular: {kpi.get('error', 'desconocido')}"

        # Diapositiva 3: productos en riesgo (tabla)
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Productos clase A en riesgo (top 10)"
        if isinstance(riesgo, list) and riesgo:
            filas, columnas = len(riesgo) + 1, 4
            tabla = slide.shapes.add_table(
                filas, columnas, Inches(0.5), Inches(1.5), Inches(12), Inches(5)
            ).table
            for c, encabezado in enumerate(["Código", "Producto", "Variación", "Ventas 2010"]):
                tabla.cell(0, c).text = encabezado
            for i, p in enumerate(riesgo, start=1):
                tabla.cell(i, 0).text = str(p["stock_code"])
                tabla.cell(i, 1).text = str(p.get("description") or "")
                tabla.cell(i, 2).text = f"{p['variacion_pct']:.1%}"
                tabla.cell(i, 3).text = f"{p['ventas_2010']:,.2f}"

        # Diapositiva 4: cancelaciones (tabla)
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Productos con más cancelaciones (top 10)"
        if isinstance(cancelacion, list) and cancelacion:
            filas, columnas = len(cancelacion) + 1, 4
            tabla = slide.shapes.add_table(
                filas, columnas, Inches(0.5), Inches(1.5), Inches(12), Inches(5)
            ).table
            for c, encabezado in enumerate(["Código", "Producto", "Unidades vendidas", "Tasa"]):
                tabla.cell(0, c).text = encabezado
            for i, p in enumerate(cancelacion, start=1):
                tabla.cell(i, 0).text = str(p["stock_code"])
                tabla.cell(i, 1).text = str(p.get("description") or "")
                tabla.cell(i, 2).text = str(p["unidades_vendidas"])
                tabla.cell(i, 3).text = f"{p['tasa']:.1%}"

        ruta = _nombre_archivo(titulo, "pptx")
        prs.save(str(ruta))
        return {
            "archivo": str(ruta.relative_to(RAIZ)),
            "mensaje": f"Presentación guardada en {ruta.relative_to(RAIZ)}",
        }
    except Exception as e:
        return {"error": str(e)}


if __name__ == "__main__":
    mcp.run(transport="stdio")
